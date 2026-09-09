"""Fixture-driven tests for the allocation extraction MVP."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from scm_allocation.extraction import extract_email, is_out_of_scope_email
from scm_allocation.models.allocation import ExtractionStatus
from scm_allocation.models.reference import AllocationContext
from scm_allocation.normalization import resolve_storage_locations
from scm_allocation.parsers import parse_html
from scm_allocation.parsers.excel import parse_excel


ROOT = Path(__file__).resolve().parents[2]
INSPECTION = ROOT / "data" / "inspection"
ATTACHMENTS = ROOT / "tests" / "fixtures" / "attachments"
EXCEL = ATTACHMENTS / "Alokasi Logitech BIC1 to iBox BCA XE05 7 September.xlsx"
if not EXCEL.exists():
    EXCEL = None


def load_email(index: int) -> dict:
    return json.loads((INSPECTION / f"email_{index:03d}.json").read_text(encoding="utf-8"))


@pytest.mark.skipif(EXCEL is None, reason="local confidential Excel fixture is unavailable")
def test_excel_sample_extracts_compound_grid():
    result = parse_excel(EXCEL, source_email_id="email-002")

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 56
    assert result.records[0].material_code == "8100090031"
    assert result.records[0].quantity == 10
    assert result.records[0].destination_plant_code == "XE05"
    assert result.records[0].issuing_warehouse_code == "BIC1"
    assert result.records[0].allocation_context is AllocationContext.ACCESSORIES
    assert all(record.source_filename == EXCEL.name for record in result.records)


@pytest.mark.skipif(not (INSPECTION / "email_003_body.html").exists(), reason="local HTML fixture is unavailable")
def test_pivot_sample_extracts_request_column():
    html = (INSPECTION / "email_003_body.html").read_text(encoding="utf-8")
    result = parse_html(html, source_email_id="email-003")

    assert len(result.records) == 18
    assert result.records[0].material_code == "8100226821"
    assert result.records[0].quantity == 20
    assert result.records[0].destination_plant_code == "X071"
    assert result.records[0].allocation_context is AllocationContext.DEVICE


@pytest.mark.skipif(not (INSPECTION / "email_004_body.html").exists(), reason="local HTML fixture is unavailable")
def test_flat_sample_extracts_both_tables():
    html = (INSPECTION / "email_004_body.html").read_text(encoding="utf-8")
    result = parse_html(html, source_email_id="email-004")

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 44
    assert result.records[0].material_code == "8100261734"
    assert result.records[0].quantity == 5
    assert result.records[0].destination_plant_code == "X115"
    assert result.records[0].issuing_warehouse_code == "BIC1"


@pytest.mark.skipif(not (INSPECTION / "email_005_body.html").exists(), reason="local HTML fixture is unavailable")
def test_second_pivot_sample_extracts_all_requests():
    html = (INSPECTION / "email_005_body.html").read_text(encoding="utf-8")
    result = parse_html(html, source_email_id="email-005")

    assert len(result.records) == 7
    assert [record.quantity for record in result.records] == [1, 2, 2, 4, 4, 3, 5]
    assert all(record.destination_plant_code == "X071" for record in result.records)


def test_malformed_flat_row_returns_partial_record():
    html = """
    <table>
      <tr><th>PLAN</th><th>plant_name</th><th>MATERIAL</th><th>BRAND TYPE 2</th><th>PO</th></tr>
      <tr><td>X001</td><td>Test Store</td><td>123</td><td>Test Item</td><td></td></tr>
    </table>
    """

    result = parse_html(html, source_email_id="synthetic")

    assert len(result.records) == 1
    assert result.status is ExtractionStatus.PARTIAL
    assert result.records[0].quantity is None
    assert "missing or invalid quantity" in result.records[0].errors


@pytest.mark.skipif(not (INSPECTION / "email_001.json").exists(), reason="local email fixture is unavailable")
def test_batam_email_is_skipped_before_attachment_parsing():
    email = load_email(1)
    result = extract_email(email)

    assert result.status is ExtractionStatus.SKIPPED
    assert result.records == []
    assert result.parser_type == "out_of_scope"


@pytest.mark.skipif(EXCEL is None, reason="local confidential Excel fixture is unavailable")
def test_orchestration_selects_excel_parser():
    email = load_email(2)
    result = extract_email(email, attachment_paths={EXCEL.name: EXCEL})

    assert result.parser_type == "excel_compound_grid"
    assert len(result.records) == 56


@pytest.mark.skipif(not (INSPECTION / "email_003.json").exists(), reason="local email fixture is unavailable")
def test_orchestration_selects_html_parser():
    result = extract_email(load_email(3))

    assert result.parser_type == "html_pivot_table"
    assert len(result.records) == 18


def test_real_request_matrix_variant_extracts_body_labeled_plant():
    email = {
        "entry_id": "gwalk-synthetic",
        "subject": "Permintaan Alokasi Barang iBox Gwalk Citraland Surabaya // X213",
        "body_plain_text": "Store Code : iBox Citraland Gwalk Surabaya\nStore Desc : X213\nReason : Indenan customer",
        "body_html": """
            <table>
              <tr><th>Material</th><th>Material description</th><th>BIC1</th><th>Request</th><th>Ket</th></tr>
              <tr><td>8100027770</td><td>IMAC 24 inch</td><td>2</td><td>1</td><td>Pengajuan PT</td></tr>
              <tr><td>8100261692</td><td>MBA 13 MDN</td><td>69</td><td>2</td><td>Pendingan Cust</td></tr>
            </table>
        """,
        "attachments": [],
    }

    result = extract_email(email)

    assert result.status is ExtractionStatus.VALID
    assert result.parser_type == "html_request_matrix"
    assert len(result.records) == 2
    assert result.records[0].destination_plant_code == "X213"
    assert result.records[0].issuing_warehouse_code == "BIC1"
    assert result.records[0].quantity == 1


def test_batam_samsung_workflow_is_excluded_before_excel_parsing():
    subject = "RE: STO ALOKASI BATAM SAMSUNG DEVICE DAN ACCS WK36 2026"
    email = {
        "entry_id": "batam-samsung-synthetic",
        "subject": subject,
        "attachments": [{"filename": "54. Alokasi Batam Samsung wk36 2026 (1).xlsx"}],
    }

    assert is_out_of_scope_email(subject) is True
    result = extract_email(email)

    assert result.status is ExtractionStatus.SKIPPED
    assert result.parser_type == "out_of_scope"
    assert result.records == []


def _add_simple_allocation_sheet(workbook, name, rows):
    worksheet = workbook.create_sheet(name)
    worksheet.append(["Material", "Material description", "Site Code", "Site Desc", "Qty"])
    for row in rows:
        worksheet.append(row)
    return worksheet


def test_excel_allocation_sheet_can_follow_irrelevant_first_sheet(tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    _add_simple_allocation_sheet(workbook, "Request", [["8100000001", "Sample Item", "X001", "Sample Store", 3]])
    path = tmp_path / "non_first.xlsx"
    workbook.save(path)

    result = parse_excel(path, source_email_id="non-first")

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 1
    assert result.records[0].source_sheet == "Request"
    assert result.records[0].material_code == "8100000001"
    assert result.metadata["worksheet_diagnostics"][0]["status"] == "skipped"


def test_excel_skips_multiple_irrelevant_sheets_before_allocation_sheet(tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Control"
    workbook.create_sheet("SAP Export")
    workbook.create_sheet("Summary")
    _add_simple_allocation_sheet(workbook, "Allocation", [["8100000002", "Another Item", "X002", "Another Store", 4]])
    path = tmp_path / "irrelevant_before.xlsx"
    workbook.save(path)

    result = parse_excel(path)

    assert len(result.records) == 1
    assert result.records[0].source_sheet == "Allocation"
    skipped = [item for item in result.metadata["worksheet_diagnostics"] if item["status"] == "skipped"]
    assert len(skipped) == 3


def test_excel_extracts_multiple_valid_worksheets(tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    _add_simple_allocation_sheet(workbook, "APPLE", [["8100000003", "Apple Item", "X003", "Apple Store", 2]])
    _add_simple_allocation_sheet(workbook, "DJI", [["8100000004", "DJI Item", "X004", "DJI Store", 5]])
    path = tmp_path / "multiple_valid.xlsx"
    workbook.save(path)

    result = parse_excel(path)

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 2
    assert {record.source_sheet for record in result.records} == {"APPLE", "DJI"}
    assert result.parser_type == "simple_excel_table"


def test_excel_skips_blank_request_rows(tmp_path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Request"
    worksheet.append(["Material", "Material description", "Request", "BIC1"])
    worksheet.append(["8100000007", "Stock-only Item", None, 12])
    worksheet.append(["8100000008", "Requested Item", 2, 12])
    path = tmp_path / "blank_requests.xlsx"
    workbook.save(path)

    result = parse_excel(path, destination_plant_code="U064")

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 1
    assert result.records[0].material_code == "8100000008"


def test_excel_uses_subject_plant_when_body_has_no_plant_label(tmp_path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Request"
    worksheet.append(["Material", "Material description", "Request", "BIC1"])
    worksheet.append(["8100000009", "Subject Plant Item", 3, 20])
    path = tmp_path / "subject_plant.xlsx"
    workbook.save(path)

    email = {
        "entry_id": "subject-plant",
        "subject": "Request allocation for UR Pakuwon City (U064)",
        "body_plain_text": "Please process the request.",
        "attachments": [{"filename": path.name}],
    }
    result = extract_email(email, attachment_paths={path.name: path})

    assert result.status is ExtractionStatus.VALID
    assert result.records[0].destination_plant_code == "U064"


def test_excel_destination_matrix_extracts_each_request_column(tmp_path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append([None, None, "U001", "U002"])
    worksheet.append(["Article Code", "Article Description", "UR Store One", "UR Store Two"])
    worksheet.append(["8100000010", "Matrix Item", 2, 3])
    path = tmp_path / "destination_matrix.xlsx"
    workbook.save(path)

    result = parse_excel(path)

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 2
    assert {record.destination_plant_code for record in result.records} == {"U001", "U002"}
    assert all(record.source_sheet == "Sheet1" for record in result.records)


def test_excel_destination_matrix_uses_one_non_destination_code_as_warehouse(tmp_path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append([None, None, "BGC1", "F233"])
    worksheet.append(["Article SAP", "Article Description SAP", "ERAFONE MDC CIKUPA", "Erafone Palu Grand Mall"])
    worksheet.append(["8100000011", "Wearable Item", 20, 3])
    path = tmp_path / "warehouse_destination_matrix.xlsx"
    workbook.save(path)

    result = parse_excel(path, destination_plant_code="F233")

    assert len(result.records) == 1
    assert result.records[0].destination_plant_code == "F233"
    assert result.records[0].issuing_warehouse_code == "BGC1"


def test_excel_npi_wide_site_material_matrix_extracts_each_material(tmp_path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append([None, None, "8100294963", "8100294964"])
    worksheet.append(["Site Code", "Site Desc", "ANK NANO CHARGER 45W - BLUE", "ANK POWERBANK 25K"])
    worksheet.append(["X015", "IBOX APP PLAZA INDONESIA", 5, 3])
    path = tmp_path / "npi_wide_matrix.xlsx"
    workbook.save(path)

    result = parse_excel(path, source_email_id="npi-email")

    assert result.status is ExtractionStatus.VALID
    assert result.parser_type == "excel_site_material_matrix"
    assert [(record.material_code, record.quantity) for record in result.records] == [("8100294963", 5), ("8100294964", 3)]
    assert {record.destination_plant_code for record in result.records} == {"X015"}
    assert {record.destination_plant_description for record in result.records} == {"IBOX APP PLAZA INDONESIA"}
    assert {record.destination_sloc for record in result.records} == {"1001"}


def test_extract_email_dispatches_npi_attachment_to_site_material_matrix(tmp_path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    material_codes = [f"81002949{index:02d}" for index in range(10)]
    worksheet.append([None, None, *material_codes])
    worksheet.append(["Site Code", "Site Desc", *[f"ANK Charger {index}" for index in range(10)]])
    worksheet.append(["X015", "IBOX APP PLAZA INDONESIA", *range(1, 11)])
    path = tmp_path / "npi_anker.xlsx"
    workbook.save(path)

    result = extract_email(
        {"entry_id": "npi-email", "subject": "Request Alokasi NPI Anker", "attachments": [{"filename": "Untitled attachment 00418.htm"}, {"filename": path.name}]},
        attachment_paths={path.name: path},
    )

    assert result.status is ExtractionStatus.VALID
    assert result.parser_type == "excel_site_material_matrix"
    assert len(result.records) == 10
    assert {record.source_email_id for record in result.records} == {"npi-email"}
    assert {record.source_filename for record in result.records} == {path.name}
    assert [item["status"] for item in result.metadata["attachment_diagnostics"]] == ["ignored", "selected"]


def test_excel_new_store_it_and_loops_preserve_explicit_source_and_destination_sloc(tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    headers = ["No", "Article Number (SAP Code)", "Material Description", "Qty", "Plant Asal", "Storage Location (Sloc)", "Plant Code", "Storage Location (Sloc)"]
    for sheet_name, material, quantity in (("IT", "IT1 USB C Cable", 10), ("LOOPS", "LOP Cable A to C", 5)):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(headers)
        worksheet.append([1, "8100013184", material, quantity, "BCC1", "1001", "F233", "1001"])
    stores = workbook.create_sheet("Stores")
    stores.append(["List New Store", None])
    stores.append(["F233", "ERAFONE PALU GRAND MALL"])
    path = tmp_path / "new_store.xlsx"
    workbook.save(path)

    result = parse_excel(path, source_email_id="new-store-email")

    assert result.status is ExtractionStatus.VALID
    assert result.parser_type == "excel_new_store_table"
    assert {record.source_sheet for record in result.records} == {"IT", "LOOPS"}
    assert {record.destination_plant_description for record in result.records} == {"ERAFONE PALU GRAND MALL"}
    assert {record.issuing_warehouse_code for record in result.records} == {"BCC1"}
    assert {record.issuing_warehouse_sloc for record in result.records} == {"1001"}
    assert {record.destination_sloc for record in result.records} == {"1001"}


def test_extract_email_dispatches_new_store_attachment_across_it_and_loops(tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    headers = ["No", "Article Number (SAP Code)", "Material Description", "Qty", "Plant Asal", "Storage Location (Sloc)", "Plant Code", "Storage Location (Sloc)"]
    for sheet_name, start in (("IT", 0), ("LOOPS", 7)):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(headers)
        for row_number in range(7):
            worksheet.append([row_number + 1, f"810000{start + row_number:04d}", f"{sheet_name} Cable {row_number}", row_number + 1, "BCC1", "1001", "F233", "1001"])
    stores = workbook.create_sheet("Stores")
    stores.append(["List New Store", None])
    stores.append(["F233", "ERAFONE PALU GRAND MALL"])
    path = tmp_path / "new_store_it_loops.xlsx"
    workbook.save(path)

    result = extract_email(
        {"entry_id": "new-store-email", "subject": "STO CC IT LOOPS Alokasi New Store", "attachments": [{"filename": "image001.gif"}, {"filename": path.name}]},
        attachment_paths={path.name: path},
    )

    assert result.status is ExtractionStatus.VALID
    assert result.parser_type == "excel_new_store_table"
    assert len(result.records) == 14
    assert {record.source_sheet for record in result.records} == {"IT", "LOOPS"}
    assert {record.destination_plant_description for record in result.records} == {"ERAFONE PALU GRAND MALL"}
    assert [item["status"] for item in result.metadata["attachment_diagnostics"]] == ["ignored", "selected"]


def test_extract_email_reports_a_malformed_excel_attachment_without_records(tmp_path: Path):
    path = tmp_path / "malformed.xlsx"
    path.write_bytes(b"not an Excel workbook")

    result = extract_email(
        {"entry_id": "bad-attachment", "subject": "Allocation request", "attachments": [{"filename": path.name}]},
        attachment_paths={path.name: path},
    )

    assert result.status is ExtractionStatus.INVALID
    assert result.records == []
    assert result.errors


@pytest.mark.parametrize("source", ["BGC1", "BHC1", "BIC1", "BFC1"])
def test_central_source_defaults_device_and_non_device_sloc_when_absent(tmp_path: Path, source: str):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["Material", "Material description", "Warehouse Code", "Site Code", "Qty"])
    worksheet.append(["8100000101", "MacBook Air", source, "X001", 1])
    worksheet.append(["8100000102", "USB C Cable", source, "X001", 2])
    path = tmp_path / f"central_{source}.xlsx"
    workbook.save(path)

    result = parse_excel(path)

    assert result.status is ExtractionStatus.VALID
    assert [record.issuing_warehouse_sloc for record in result.records] == ["1005", "1001"]
    assert {record.destination_sloc for record in result.records} == {"1001"}


def test_explicit_storage_locations_override_central_and_destination_defaults():
    source_sloc, destination_sloc = resolve_storage_locations(
        issuing_warehouse_code="BIC1",
        allocation_context=AllocationContext.DEVICE,
        explicit_issuing_sloc="1999",
        explicit_destination_sloc="1888",
        destination_plant_code="RDC1",
    )

    assert source_sloc == "1999"
    assert destination_sloc == "1888"


def test_html_code_store_request_aliases_are_supported():
    html = """
        <table>
          <tr><th>Division</th><th>Code Store</th><th>Store</th><th>Material</th><th>Material Desc.</th><th>Qty Request</th><th>Ket</th></tr>
          <tr><td>ED SALES REGION 4</td><td>X210</td><td>iBox Roxy Square</td><td>8100000012</td><td>Mac Item</td><td>2</td><td>Stock Kosong</td></tr>
        </table>
    """

    result = parse_html(html)

    assert result.status is ExtractionStatus.PARTIAL
    assert len(result.records) == 1
    assert result.records[0].destination_plant_code == "X210"
    assert result.records[0].quantity == 2


def test_extract_email_merges_supported_records_from_multiple_workbooks(tmp_path: Path):
    first = openpyxl.Workbook()
    _add_simple_allocation_sheet(first, "Request", [["8100000013", "First Item", "X013", "First Store", 1]])
    first_path = tmp_path / "allocation.xlsx"
    first.save(first_path)

    second = openpyxl.Workbook()
    _add_simple_allocation_sheet(second, "Stock", [["8100000014", "Second Item", "X014", "Second Store", 2]])
    second_path = tmp_path / "reference.xlsx"
    second.save(second_path)

    email = {
        "entry_id": "multi-workbook",
        "subject": "Allocation workbook pair",
        "attachments": [
            {"filename": first_path.name},
            {"filename": second_path.name},
        ],
    }
    result = extract_email(email, attachment_paths={first_path.name: first_path, second_path.name: second_path})

    assert result.status is ExtractionStatus.VALID
    assert {record.material_code for record in result.records} == {"8100000013", "8100000014"}
    assert {record.source_filename for record in result.records} == {first_path.name, second_path.name}


def test_excel_site_allocation_without_material_identity_is_review(tmp_path: Path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "allocation"
    worksheet.append(["Site Code", "Site Desc", "Allocation"])
    worksheet.append(["E894", "ERAFONE RUKO SUDIRMAN SUKOHARJO", 4])
    workbook.create_sheet("Sheet2")["A1"] = "site"
    path = tmp_path / "site_only_allocation.xlsx"
    workbook.save(path)

    result = parse_excel(path)

    assert result.status is ExtractionStatus.PARTIAL
    assert len(result.records) == 1
    assert result.records[0].material_code is None
    assert result.records[0].destination_plant_code == "E894"
    assert "no material code or material description" in result.records[0].errors[0]
    assert result.metadata["worksheet_diagnostics"][1]["status"] == "skipped"


def test_excel_deduplicates_records_across_valid_worksheets(tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    row = ["8100000005", "Duplicate Item", "X005", "Duplicate Store", 1]
    _add_simple_allocation_sheet(workbook, "Request A", [row])
    _add_simple_allocation_sheet(workbook, "Request B", [row])
    path = tmp_path / "duplicate_tabs.xlsx"
    workbook.save(path)

    result = parse_excel(path)

    assert len(result.records) == 1
    assert result.records[0].source_sheet == "Request A"
    assert result.metadata["worksheet_diagnostics"][1]["duplicates_skipped"] == 1


def test_excel_reports_no_supported_structure_after_inspecting_all_sheets(tmp_path: Path):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    workbook.create_sheet("SAP Export")
    workbook["Summary"]["A1"] = "Grand Total"
    workbook["SAP Export"]["A1"] = "Material Document"
    path = tmp_path / "no_allocation_tabs.xlsx"
    workbook.save(path)

    result = parse_excel(path)

    assert result.status is ExtractionStatus.INVALID
    assert result.errors == ("no allocation worksheet found",)
    assert len(result.metadata["worksheet_diagnostics"]) == 2
    assert all(item["status"] == "skipped" for item in result.metadata["worksheet_diagnostics"])


def test_html_headers_with_blank_cells_return_controlled_status():
        html = """
                <table>
                    <tr><th>MATERIAL</th><th>Row Labels</th><th></th><th>Request</th></tr>
                    <tr><th>BIC1</th><th></th><th>BCC1</th><th>X001</th></tr>
                    <tr><td>8100000006</td><td>Item</td><td>2</td><td>1</td></tr>
                </table>
        """

        result = parse_html(html)

        assert result.status in {ExtractionStatus.INVALID, ExtractionStatus.PARTIAL}


def test_attachment_order_independence_and_non_xlsx_filtering(tmp_path: Path):
    workbook = openpyxl.Workbook()
    _add_simple_allocation_sheet(workbook, "Request", [["8100000015", "Order Test Item", "X015", "Order Store", 5]])
    xlsx_path = tmp_path / "allocation.xlsx"
    workbook.save(xlsx_path)

    email = {
        "entry_id": "order-test",
        "subject": "Request Allocation with multiple asset attachments",
        "attachments": [
            {"filename": "image001.gif"},
            {"filename": "company_logo.png"},
            {"filename": "signature.htm"},
            {"filename": xlsx_path.name},
            {"filename": "banner.jpg"},
        ],
    }

    result = extract_email(email, attachment_paths={xlsx_path.name: xlsx_path})

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 1
    assert result.records[0].material_code == "8100000015"
    assert [diag["status"] for diag in result.metadata["attachment_diagnostics"]] == [
        "ignored", "ignored", "ignored", "selected", "ignored"
    ]


def test_multiple_xlsx_with_unrelated_workbook_safety(tmp_path: Path):
    valid_wb = openpyxl.Workbook()
    _add_simple_allocation_sheet(valid_wb, "Request", [["8100000016", "Valid Item", "X016", "Valid Store", 10]])
    valid_path = tmp_path / "allocation.xlsx"
    valid_wb.save(valid_path)

    unrelated_wb = openpyxl.Workbook()
    sheet = unrelated_wb.active
    sheet.title = "Instructions"
    sheet.append(["Policy Name", "Description"])
    sheet.append(["SCM Guideline", "Standard operating procedure"])
    unrelated_path = tmp_path / "instructions.xlsx"
    unrelated_wb.save(unrelated_path)

    email = {
        "entry_id": "multi-xlsx-unrelated",
        "subject": "Allocation and Instructions",
        "attachments": [
            {"filename": valid_path.name},
            {"filename": unrelated_path.name},
        ],
    }

    result = extract_email(
        email,
        attachment_paths={valid_path.name: valid_path, unrelated_path.name: unrelated_path},
    )

    assert result.status is ExtractionStatus.VALID
    assert len(result.records) == 1
    assert result.records[0].material_code == "8100000016"
