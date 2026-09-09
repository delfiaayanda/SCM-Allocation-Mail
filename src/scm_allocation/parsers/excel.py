"""Parser for allocation worksheets in multi-sheet Excel workbooks."""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any, Optional

import openpyxl

from scm_allocation.models.allocation import AllocationRecord, ExtractionResult, ExtractionStatus
from scm_allocation.normalization import find_alias_index, normalize_quantity, normalize_text
from scm_allocation.parsers.common import infer_allocation_context, make_record


def _category_bands(ws: Any, row_number: int) -> list[Optional[str]]:
    bands: list[Optional[str]] = []
    current: Optional[str] = None
    for column in range(1, ws.max_column + 1):
        value = normalize_text(ws.cell(row_number, column).value)
        if value:
            current = value.casefold()
        bands.append(current)
    return bands


def _headers(ws: Any, row_number: int) -> list[Optional[str]]:
    return [normalize_text(ws.cell(row_number, column).value) for column in range(1, ws.max_column + 1)]


def _find_compound_header(ws: Any) -> Optional[int]:
    for row_number in range(1, min(ws.max_row, 20) + 1):
        headers = _headers(ws, row_number)
        if any(value and value.casefold() == "article code" for value in headers):
            return row_number
    return None


def _find_simple_header(ws: Any) -> Optional[int]:
    for row_number in range(1, min(ws.max_row, 20) + 1):
        headers = _headers(ws, row_number)
        if (
            find_alias_index(headers, "material_code") is not None
            and find_alias_index(headers, "material_description") is not None
            and find_alias_index(headers, "quantity") is not None
        ):
            return row_number
    return None


def _find_site_allocation_header(ws: Any) -> Optional[int]:
    for row_number in range(1, min(ws.max_row, 20) + 1):
        headers = _headers(ws, row_number)
        header_names = {value.casefold() for value in headers if value}
        if {"site code", "site desc", "allocation"}.issubset(header_names):
            return row_number
    return None


def _find_text_index(headers: list[Optional[str]], *terms: str) -> Optional[int]:
    normalized_terms = tuple(term.casefold() for term in terms)
    for index, header in enumerate(headers):
        value = (header or "").casefold()
        if any(term in value for term in normalized_terms):
            return index
    return None


def _compound_columns(ws: Any, header_row: int) -> Optional[dict[str, int]]:
    headers = _headers(ws, header_row)
    bands = _category_bands(ws, max(1, header_row - 1))
    columns: dict[str, Optional[int]] = {
        "material_code": next((i for i, value in enumerate(headers) if value and value.casefold() == "article code"), None),
        "material_description": next((i for i, value in enumerate(headers) if value and value.casefold() == "article description"), None),
        "quantity": next((i for i, value in enumerate(headers) if value and value.casefold() == "alokasi"), None),
    }
    for field_name, band_name, header_name in (
        ("from_code", "from", "site code"),
        ("from_description", "from", "site desc"),
        ("to_code", "to store", "site code"),
        ("to_description", "to store", "site desc"),
    ):
        columns[field_name] = next(
            (i for i, value in enumerate(headers) if value and value.casefold() == header_name and bands[i] == band_name),
            None,
        )
    if any(index is None for index in columns.values()):
        return None
    return {name: index for name, index in columns.items() if index is not None}


def _simple_columns(ws: Any, header_row: int) -> Optional[dict[str, Optional[int]]]:
    headers = _headers(ws, header_row)
    material_index = find_alias_index(headers, "material_code")
    description_index = find_alias_index(headers, "material_description")
    quantity_index = find_alias_index(headers, "quantity")
    if material_index is None or description_index is None or quantity_index is None:
        return None

    marker_headers = {value.casefold() for value in headers if value}
    has_allocation_marker = bool(
        {"request", "site code", "site desc", "from", "wh code", "warehouse code"} & marker_headers
        or any(value and (value.casefold().startswith("gudang") or value.casefold().startswith("ur ")) for value in headers)
    )
    if not has_allocation_marker:
        return None

    destination_code = find_alias_index(headers, "destination_plant_code")
    destination_description = find_alias_index(headers, "destination_plant_description")
    issuing_code = find_alias_index(headers, "issuing_warehouse_code")
    issuing_description = find_alias_index(headers, "issuing_warehouse_description")
    if destination_code is None:
        destination_code = next(
            (index for index, value in enumerate(headers) if value and value.casefold() == "wh code"),
            None,
        )
    if destination_description is None:
        destination_description = _find_text_index(headers, "bu desc", "plant desc", "ur ")
    if issuing_code is None:
        issuing_code = _find_text_index(headers, "from site code", "from", "warehouse code")
    if issuing_description is None:
        issuing_description = _find_text_index(headers, "from site desc", "gudang")

    return {
        "material_code": material_index,
        "material_description": description_index,
        "quantity": quantity_index,
        "destination_code": destination_code,
        "destination_description": destination_description,
        "issuing_code": issuing_code,
        "issuing_description": issuing_description,
    }


def _find_destination_matrix(ws: Any) -> Optional[dict[str, Any]]:
    """Detect a wide allocation matrix with destination codes above columns."""
    for header_row in range(2, min(ws.max_row, 20) + 1):
        headers = _headers(ws, header_row)
        material_index = find_alias_index(headers, "material_code")
        description_index = find_alias_index(headers, "material_description")
        if material_index is None or description_index is None:
            continue
        code_row = header_row - 1
        code_columns = {
            column - 1: normalize_text(ws.cell(code_row, column).value)
            for column in range(1, ws.max_column + 1)
            if normalize_text(ws.cell(code_row, column).value)
            and re.fullmatch(r"[A-Za-z]{1,5}\d{1,4}", normalize_text(ws.cell(code_row, column).value) or "")
        }
        if not code_columns:
            continue
        has_quantity = any(
            normalize_quantity(ws.cell(row, column + 1).value) is not None
            for row in range(header_row + 1, min(ws.max_row, header_row + 20) + 1)
            for column in code_columns
        )
        if has_quantity:
            return {
                "header_row": header_row,
                "material_code": material_index,
                "material_description": description_index,
                "code_columns": code_columns,
            }
    return None


def _find_new_store_header(ws: Any) -> Optional[int]:
    for row_number in range(1, min(ws.max_row, 20) + 1):
        headers = _headers(ws, row_number)
        names = {(header or "").casefold() for header in headers}
        if {"article number (sap code)", "material description", "qty", "plant asal", "plant code"}.issubset(names):
            return row_number
    return None


def _find_site_material_matrix(ws: Any) -> Optional[dict[str, Any]]:
    """Detect a wide matrix: material codes row 1, descriptions row 2, sites below."""
    if ws.max_row < 3 or ws.max_column < 3:
        return None
    headers = _headers(ws, 2)
    if not (headers[0] and headers[0].casefold() == "site code" and headers[1] and headers[1].casefold() == "site desc"):
        return None
    columns = {
        index: normalize_text(ws.cell(1, index + 1).value)
        for index in range(2, ws.max_column)
        if normalize_text(ws.cell(1, index + 1).value)
        and re.fullmatch(r"\d{6,}", normalize_text(ws.cell(1, index + 1).value) or "")
        and headers[index]
    }
    return {"header_row": 2, "material_columns": columns} if columns else None


def _site_material_matrix_records(
    worksheet: Any,
    matrix: dict[str, Any],
    *,
    filepath: Path,
    source_email_id: Optional[str],
) -> list[AllocationRecord]:
    records: list[AllocationRecord] = []
    headers = _headers(worksheet, matrix["header_row"])
    for row_number in range(matrix["header_row"] + 1, worksheet.max_row + 1):
        values = _row_values(worksheet, row_number)
        destination_code = values[0] if values else None
        if _is_total_or_empty(destination_code):
            continue
        for column, material_code in matrix["material_columns"].items():
            quantity = values[column] if column < len(values) else None
            if normalize_quantity(quantity) in {None, 0}:
                continue
            records.append(make_record(
                material_code=material_code,
                material_description=headers[column],
                quantity=quantity,
                destination_plant_code=destination_code,
                destination_plant_description=values[1] if len(values) > 1 else None,
                issuing_warehouse_code=None,
                issuing_warehouse_description=None,
                allocation_context=infer_allocation_context(worksheet.title, headers[column]),
                source_email_id=source_email_id,
                source_filename=filepath.name,
                source_sheet=worksheet.title,
                source_row=row_number,
                source_location=f"{worksheet.title}!{column + 1}{row_number}",
                parser_type="excel_site_material_matrix",
                confidence=0.95,
            ))
    return records


def _new_store_columns(ws: Any, header_row: int) -> Optional[dict[str, int]]:
    """Detect the IT/LOOPS new-store table with source and destination SLoc columns."""
    headers = _headers(ws, header_row)
    def exact(value: str) -> Optional[int]:
        return next((index for index, header in enumerate(headers) if header and header.casefold() == value), None)
    material = exact("article number (sap code)")
    description = exact("material description")
    quantity = exact("qty")
    source = exact("plant asal")
    destination = exact("plant code")
    if None in {material, description, quantity, source, destination}:
        return None
    source_sloc = source + 1 if source + 1 < len(headers) and headers[source + 1] and "storage location" in headers[source + 1].casefold() else None
    destination_sloc = destination + 1 if destination + 1 < len(headers) and headers[destination + 1] and "storage location" in headers[destination + 1].casefold() else None
    if source_sloc is None or destination_sloc is None:
        return None
    return {"material_code": material, "material_description": description, "quantity": quantity, "issuing_code": source, "destination_code": destination, "issuing_sloc": source_sloc, "destination_sloc": destination_sloc}


def _new_store_records(
    worksheet: Any,
    header_row: int,
    columns: dict[str, int],
    *,
    filepath: Path,
    source_email_id: Optional[str],
    destination_names: dict[str, str],
) -> list[AllocationRecord]:
    records: list[AllocationRecord] = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        values = _row_values(worksheet, row_number)
        material = values[columns["material_code"]] if columns["material_code"] < len(values) else None
        if _is_total_or_empty(material):
            continue
        destination = values[columns["destination_code"]] if columns["destination_code"] < len(values) else None
        records.append(make_record(
            material_code=material,
            material_description=values[columns["material_description"]],
            quantity=values[columns["quantity"]],
            destination_plant_code=destination,
            destination_plant_description=destination_names.get(normalize_text(destination) or ""),
            issuing_warehouse_code=values[columns["issuing_code"]],
            issuing_warehouse_description=None,
            issuing_warehouse_sloc=values[columns["issuing_sloc"]],
            destination_sloc=values[columns["destination_sloc"]],
            allocation_context=infer_allocation_context(worksheet.title, values[columns["material_description"]]),
            source_email_id=source_email_id,
            source_filename=filepath.name,
            source_sheet=worksheet.title,
            source_row=row_number,
            source_location=f"{worksheet.title}!A{row_number}",
            parser_type="excel_new_store_table",
            confidence=0.95,
        ))
    return records


def _new_store_destination_names(workbook: Any) -> dict[str, str]:
    names: dict[str, str] = {}
    for worksheet in workbook.worksheets:
        if (normalize_text(worksheet.cell(1, 1).value) or "").casefold() != "list new store":
            continue
        for row_number in range(2, worksheet.max_row + 1):
            code = normalize_text(worksheet.cell(row_number, 1).value)
            description = normalize_text(worksheet.cell(row_number, 2).value)
            if code and description:
                names[code] = description
    return names


def _row_values(ws: Any, row_number: int) -> list[Any]:
    return [ws.cell(row_number, column).value for column in range(1, ws.max_column + 1)]


def _is_total_or_empty(value: Any) -> bool:
    text = normalize_text(value)
    return text is None or text.casefold() in {"grand total", "total"}


def _site_allocation_records(
    worksheet: Any,
    header_row: int,
    *,
    filepath: Path,
    source_email_id: Optional[str],
) -> list[AllocationRecord]:
    headers = _headers(worksheet, header_row)
    site_code_index = next(index for index, value in enumerate(headers) if value and value.casefold() == "site code")
    site_desc_index = next(index for index, value in enumerate(headers) if value and value.casefold() == "site desc")
    quantity_index = next(index for index, value in enumerate(headers) if value and value.casefold() == "allocation")
    records: list[AllocationRecord] = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        values = _row_values(worksheet, row_number)
        quantity = normalize_quantity(values[quantity_index] if quantity_index < len(values) else None)
        if quantity is None or quantity == 0:
            continue
        records.append(
            make_record(
                material_code=None,
                material_description=None,
                quantity=quantity,
                destination_plant_code=values[site_code_index] if site_code_index < len(values) else None,
                destination_plant_description=values[site_desc_index] if site_desc_index < len(values) else None,
                issuing_warehouse_code=None,
                issuing_warehouse_description=None,
                allocation_context=infer_allocation_context(worksheet.title),
                source_email_id=source_email_id,
                source_filename=filepath.name,
                source_sheet=worksheet.title,
                source_row=row_number,
                source_location=f"{worksheet.title}!A{row_number}",
                parser_type="excel_site_allocation",
                confidence=0.7,
                extra_errors=("source has no material code or material description",),
            )
        )
    return records


def _matrix_records(
    worksheet: Any,
    matrix: dict[str, Any],
    *,
    filepath: Path,
    source_email_id: Optional[str],
    destination_plant_code: Optional[str],
    destination_plant_description: Optional[str],
) -> list[AllocationRecord]:
    header_row = matrix["header_row"]
    material_index = matrix["material_code"]
    description_index = matrix["material_description"]
    code_columns: dict[int, str] = matrix["code_columns"]
    destination_codes = set(code_columns.values())
    if destination_plant_code and destination_plant_code in destination_codes:
        selected_destinations = {destination_plant_code}
        source_candidates = [code for code in code_columns.values() if code != destination_plant_code]
        issuing_warehouse_code = source_candidates[0] if len(source_candidates) == 1 else None
    else:
        selected_destinations = destination_codes
        issuing_warehouse_code = None

    records: list[AllocationRecord] = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        values = _row_values(worksheet, row_number)
        material_value = values[material_index] if material_index < len(values) else None
        if _is_total_or_empty(material_value):
            continue
        description = values[description_index] if description_index < len(values) else None
        for column, code in code_columns.items():
            if code not in selected_destinations:
                continue
            quantity = normalize_quantity(values[column] if column < len(values) else None)
            if quantity is None or quantity == 0:
                continue
            headers = _headers(worksheet, header_row)
            record = make_record(
                material_code=material_value,
                material_description=description,
                quantity=quantity,
                destination_plant_code=code,
                destination_plant_description=headers[column] or destination_plant_description,
                issuing_warehouse_code=issuing_warehouse_code,
                issuing_warehouse_description=None,
                allocation_context=infer_allocation_context(worksheet.title, description),
                source_email_id=source_email_id,
                source_filename=filepath.name,
                source_sheet=worksheet.title,
                source_row=row_number,
                source_location=f"{worksheet.title}!{column + 1}{row_number}",
                parser_type="excel_destination_matrix",
                confidence=0.9,
            )
            records.append(record)
    return records


def parse_excel(
    filepath: Path,
    *,
    source_email_id: Optional[str] = None,
    destination_plant_code: Optional[str] = None,
    destination_plant_description: Optional[str] = None,
) -> ExtractionResult:
    """Inspect every worksheet and extract records from all supported allocation sheets."""
    parser_type = "excel_compound_grid"
    try:
        workbook = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as err:
        return ExtractionResult([], ExtractionStatus.INVALID, parser_type, source_email_id, (str(err),))

    records: list[AllocationRecord] = []
    seen_records: set[tuple[object, ...]] = set()
    parser_types: set[str] = set()
    diagnostics: list[dict[str, object]] = []
    matched_sheet_count = 0
    destination_names = _new_store_destination_names(workbook)

    for worksheet in workbook.worksheets:
        compound_header = _find_compound_header(worksheet)
        compound_columns = _compound_columns(worksheet, compound_header) if compound_header else None
        matrix = None if compound_columns else _find_destination_matrix(worksheet)
        site_material_matrix = None if compound_columns or matrix else _find_site_material_matrix(worksheet)
        site_allocation_header = None if compound_columns or matrix or site_material_matrix else _find_site_allocation_header(worksheet)
        new_store_header = None if compound_columns or matrix or site_material_matrix or site_allocation_header else _find_new_store_header(worksheet)
        simple_header = None if new_store_header is not None or compound_columns or matrix or site_material_matrix or site_allocation_header else _find_simple_header(worksheet)
        simple_columns = _simple_columns(worksheet, simple_header) if simple_header is not None else None
        new_store_columns = _new_store_columns(worksheet, new_store_header) if new_store_header is not None else None

        if compound_columns and compound_header is not None:
            format_name = "excel_compound_grid"
            header_row = compound_header
            columns: dict[str, Optional[int]] = compound_columns
        elif matrix:
            format_name = "excel_destination_matrix"
            header_row = matrix["header_row"]
            matrix_sheet_records = _matrix_records(
                worksheet,
                matrix,
                filepath=filepath,
                source_email_id=source_email_id,
                destination_plant_code=destination_plant_code,
                destination_plant_description=destination_plant_description,
            )
            matched_sheet_count += 1
            parser_types.add(format_name)
            sheet_records = 0
            duplicate_records = 0
            for record in matrix_sheet_records:
                identity = (
                    record.material_code,
                    record.material_description,
                    record.quantity,
                    record.destination_plant_code,
                    record.issuing_warehouse_code,
                )
                if identity in seen_records:
                    duplicate_records += 1
                    continue
                seen_records.add(identity)
                records.append(record)
                sheet_records += 1
            headers = _headers(worksheet, header_row)
            diagnostics.append({
                "workbook": filepath.name,
                "sheet": worksheet.title,
                "status": "matched",
                "format": format_name,
                "header_row": header_row,
                "dimensions": worksheet.dimensions,
                "headers": headers[:16],
                "sample_rows": [_row_values(worksheet, row)[:16] for row in range(max(1, header_row - 1), min(header_row + 2, worksheet.max_row) + 1)],
                "records": sheet_records,
                "duplicates_skipped": duplicate_records,
            })
            continue
        elif site_material_matrix:
            format_name = "excel_site_material_matrix"
            header_row = site_material_matrix["header_row"]
            source_records = _site_material_matrix_records(
                worksheet, site_material_matrix, filepath=filepath, source_email_id=source_email_id
            )
        elif new_store_columns and new_store_header is not None:
            format_name = "excel_new_store_table"
            header_row = new_store_header
            source_records = _new_store_records(
                worksheet, header_row, new_store_columns, filepath=filepath,
                source_email_id=source_email_id, destination_names=destination_names,
            )
        elif site_allocation_header is not None:
            format_name = "excel_site_allocation"
            header_row = site_allocation_header
            matched_sheet_count += 1
            parser_types.add(format_name)
            site_records = _site_allocation_records(
                worksheet,
                header_row,
                filepath=filepath,
                source_email_id=source_email_id,
            )
            sheet_records = 0
            duplicate_records = 0
            for record in site_records:
                identity = (
                    record.material_code,
                    record.material_description,
                    record.quantity,
                    record.destination_plant_code,
                    record.issuing_warehouse_code,
                )
                if identity in seen_records:
                    duplicate_records += 1
                    continue
                seen_records.add(identity)
                records.append(record)
                sheet_records += 1
            diagnostics.append({
                "workbook": filepath.name,
                "sheet": worksheet.title,
                "status": "matched" if sheet_records else "matched_empty",
                "format": format_name,
                "header_row": header_row,
                "dimensions": worksheet.dimensions,
                "headers": _headers(worksheet, header_row)[:16],
                "sample_rows": [_row_values(worksheet, row)[:16] for row in range(header_row, min(header_row + 3, worksheet.max_row) + 1)],
                "records": sheet_records,
                "duplicates_skipped": duplicate_records,
                "reason": "material identity missing from source" if sheet_records else "no nonblank allocation quantities",
            })
            continue
        elif simple_columns and simple_header is not None:
            format_name = "simple_excel_table"
            header_row = simple_header
            columns = simple_columns
        else:
            diagnostics.append({
                "workbook": filepath.name,
                "sheet": worksheet.title,
                "status": "skipped",
                "reason": "no supported allocation structure",
                "dimensions": worksheet.dimensions,
                    "headers": _headers(worksheet, 1)[:16] if worksheet.max_row else [],
                    "sample_rows": [_row_values(worksheet, row)[:16] for row in range(1, min(4, worksheet.max_row) + 1)],
            })
            continue

        if site_material_matrix or new_store_columns:
            matched_sheet_count += 1
            parser_types.add(format_name)
            sheet_records = 0
            duplicate_records = 0
            for record in source_records:
                identity = (record.material_code, record.material_description, record.quantity, record.destination_plant_code, record.issuing_warehouse_code)
                if identity in seen_records:
                    duplicate_records += 1
                    continue
                seen_records.add(identity)
                records.append(record)
                sheet_records += 1
            diagnostics.append({
                "workbook": filepath.name, "sheet": worksheet.title,
                "status": "matched" if sheet_records else "matched_empty", "format": format_name,
                "header_row": header_row, "dimensions": worksheet.dimensions,
                "headers": _headers(worksheet, header_row)[:16],
                "sample_rows": [_row_values(worksheet, row)[:16] for row in range(max(1, header_row - 1), min(header_row + 2, worksheet.max_row) + 1)],
                "records": sheet_records, "duplicates_skipped": duplicate_records,
                "reason": "no nonblank allocation quantities" if not sheet_records else None,
            })
            continue

        matched_sheet_count += 1
        parser_types.add(format_name)
        sheet_records = 0
        duplicate_records = 0
        for row_number in range(header_row + 1, worksheet.max_row + 1):
            values = _row_values(worksheet, row_number)
            material_index = columns["material_code"]
            material_value = values[material_index] if isinstance(material_index, int) and material_index < len(values) else None
            if _is_total_or_empty(material_value):
                continue

            description_index = columns["material_description"]
            quantity_index = columns["quantity"]
            quantity_header = _headers(worksheet, header_row)[quantity_index]
            if (
                quantity_header
                and quantity_header.casefold() in {"request", "alokasi"}
                and normalize_quantity(values[quantity_index]) is None
            ):
                continue
            destination_index = columns.get("to_code", columns.get("destination_code"))
            destination_description_index = columns.get("to_description", columns.get("destination_description"))
            issuing_index = columns.get("from_code", columns.get("issuing_code"))
            issuing_description_index = columns.get("from_description", columns.get("issuing_description"))
            row_destination = values[destination_index] if isinstance(destination_index, int) and destination_index < len(values) else destination_plant_code
            row_destination_description = (
                values[destination_description_index]
                if isinstance(destination_description_index, int) and destination_description_index < len(values)
                else destination_plant_description
            )
            row_issuing = values[issuing_index] if isinstance(issuing_index, int) and issuing_index < len(values) else None
            row_issuing_description = (
                values[issuing_description_index]
                if isinstance(issuing_description_index, int) and issuing_description_index < len(values)
                else None
            )
            description = values[description_index] if isinstance(description_index, int) and description_index < len(values) else None
            quantity = values[quantity_index] if isinstance(quantity_index, int) and quantity_index < len(values) else None
            record = make_record(
                material_code=material_value,
                material_description=description,
                quantity=quantity,
                destination_plant_code=row_destination,
                destination_plant_description=row_destination_description,
                issuing_warehouse_code=row_issuing,
                issuing_warehouse_description=row_issuing_description,
                allocation_context=infer_allocation_context(
                    values[0] if format_name == "excel_compound_grid" and values else worksheet.title,
                    description,
                ),
                source_email_id=source_email_id,
                source_filename=filepath.name,
                source_sheet=worksheet.title,
                source_row=row_number,
                source_location=f"{worksheet.title}!A{row_number}",
                parser_type=format_name,
                    confidence=1.0 if format_name == "excel_compound_grid" else 0.9,
            )
            identity = (
                record.material_code,
                record.material_description,
                record.quantity,
                record.destination_plant_code,
                record.issuing_warehouse_code,
            )
            if identity in seen_records:
                duplicate_records += 1
                continue
            seen_records.add(identity)
            records.append(record)
            sheet_records += 1

        diagnostics.append({
            "workbook": filepath.name,
            "sheet": worksheet.title,
            "status": "matched" if sheet_records else "matched_empty",
            "format": format_name,
            "header_row": header_row,
            "dimensions": worksheet.dimensions,
                "headers": _headers(worksheet, header_row)[:16],
                "sample_rows": [_row_values(worksheet, row)[:16] for row in range(header_row, min(header_row + 3, worksheet.max_row) + 1)],
            "records": sheet_records,
            "duplicates_skipped": duplicate_records,
            "reason": "no nonblank request quantities" if not sheet_records else None,
        })

    metadata = {"worksheet_diagnostics": diagnostics}
    if not records:
        errors = (
            ("allocation worksheet(s) found but no records extracted",)
            if matched_sheet_count
            else ("no allocation worksheet found",)
        )
        return ExtractionResult(
            [],
            ExtractionStatus.INVALID,
            "+".join(sorted(parser_types)) or parser_type,
            source_email_id,
            errors,
            metadata=metadata,
        )

    status = ExtractionStatus.PARTIAL if any(record.extraction_status != ExtractionStatus.VALID for record in records) else ExtractionStatus.VALID
    result_parser_type = next(iter(parser_types)) if len(parser_types) == 1 else "excel_multi_sheet"
    return ExtractionResult(records, status, result_parser_type, source_email_id, metadata=metadata)
